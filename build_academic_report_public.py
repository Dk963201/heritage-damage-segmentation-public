from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile
from lxml import etree
root=Path(__file__).parent
d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.8)
style=d.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11);style.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:
    d.styles[name].font.name='Calibri';d.styles[name].font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def page():d.add_page_break()
d.add_heading('Historical citadel damage segmentation',0)
p('Historical research progress report\nPublic edition of the 18 October 2026 evidence snapshot')
p('Historical scope: this report preserves the earlier public-data experiments. It predates later reviewed target-label development and the four-part software update. Consult docs/PROJECT_SCOPE.md and REVIEWED_IMPROVEMENT_RESULTS.md for current evidence; scores from different datasets are not interchangeable.')
h('Abstract')
p('This project investigates automated mapping of visible masonry deterioration in historical structures in Erbil and Amedi. The intended task distinguishes healthy masonry, structural cracks, surface erosion and vegetation overgrowth. Eight hundred candidate image versions have been downloaded and integrity checked. A supervised binary crack pilot achieved 73.25% foreground Dice and 57.79% foreground intersection over union on a repeatedly used UK development set. These results demonstrate an executable pipeline, but do not establish four-class performance on Kurdish heritage photographs.')
p('The principal unresolved requirement is reliable target annotation. None of the 800 candidates currently has an approved four-class ground-truth mask. A complete training implementation and annotation validation workflow are supplied, while final target evaluation and severity classification remain pending. This report therefore documents a research prototype rather than a completed four-class study.')
h('Problem and social objective')
p('Visible crack and surface-damage mapping can help conservation staff prioritize closer inspection and organize photographic records. The proposed beneficiary is the High Commission for Erbil Citadel Revitalization, HCECR. No organizational agreement or operational adoption is claimed. Image predictions cannot determine structural stability, hidden damage or repair priority without professional assessment.')
h('Research questions')
p('Can transfer learning support useful crack segmentation with limited masonry labels? Once reviewed target masks exist, which of U-Net with EfficientNet, DeepLabV3+ and FPN best separates the four required classes on held-out structures? How do site, material, image scale and lighting affect errors? Severity classification requires a separate expert-defined outcome and cannot be inferred reliably from the fraction of visibly damaged pixels alone.')
page();h('Dataset provenance and annotation')
p('The target collection contains 800 unique file hashes: 609 image versions attributed to Erbil and 191 to Amedi. Download integrity is not proof of unique scenes, usable masonry, correct site attribution or sufficient defect resolution. Wide landscapes and context photographs occur in the candidate set. Six similarity pairs were flagged for manual review. Source URLs, authors and licenses are retained in the manifests; release photos-v0.1-candidates contains ten archives of 80 candidates each.')
p('Preprocessing indexed 141122 possible 256-pixel patches across the candidates. This is a geometry count, not a count of annotated training examples. A first annotation review pack contains 37 Erbil candidates: 29 full-resolution review candidates, one exposure hold and seven context-only entries. These are triage decisions, not completed expert annotations.')
p('The annotation convention is 0 healthy masonry, 1 structural crack, 2 surface erosion, 3 vegetation overgrowth and 255 unknown or outside scope. Annotators must distinguish erosion from ordinary rough stone, open joints and missing render, and vegetation on masonry from foreground plants. Crack appearance alone does not establish structural significance. Generated masks may assist review but cannot substitute for ground truth.')
h('Auxiliary supervised sources')
p('Storey and Snowball provide 100 original UK masonry images with binary crack masks under CC BY 4.0 [1]. The development split uses 80 images for training and 20 for validation, with fixed native-resolution 256-pixel center crops. The original validation set has nine positive crops and 36528 crack pixels out of 1310720 total pixels. Background-only accuracy is already 97.21%, illustrating why pixel accuracy is inadequate by itself.')
p('A public subset from Dais and colleagues supplies 240 image-mask pairs at 224 pixels [2]. The source repository commit is pinned to 83c371c70a27852f22ab3386269c6058a203ef38. Git blob hashes and SHA256 provenance were recorded. The repository carries GPL-3.0; a separate dataset license was not found. Raw public source images are not redistributed in this project update. These binary masks do not label erosion, vegetation or healthy masonry.')
page();h('Methods and implementation')
p('The four-class implementation uses the same reviewed source split for U-Net, DeepLabV3+ and FPN, each with EfficientNet-B0. Proposed defaults are 256-pixel patches, ImageNet encoder initialization, batch size two, 30 epochs, Adam at 0.0001 and a five-epoch linear learning-rate warmup. These are explicit engineering choices; teacher-prescribed patch dimensions were not recoverable from the extracted assignment text.')
p('Sources are assigned to splits by building before patch extraction. The manifest validator rejects repeated exact source hashes, cross-split building groups, missing review identities, invalid mask dimensions or class IDs, empty splits and training data missing any required class. The validator cannot certify annotation accuracy or authenticate a claimed human review. Near-duplicate and scene grouping still require review.')
p('Training applies CLAHE, paired horizontal and vertical flips, and brightness/color jitter to the image only. The objective combines multiclass focal loss with gamma two and soft Dice. Unknown pixels are excluded from both terms. Padded mask regions are unknown. All-unknown batches are skipped. Validation averages overlapping tile logits to reconstruct each source and counts each annotated pixel once.')
p('Checkpoint selection uses validation mean IoU. Test evaluation is disabled by default and should occur only after freezing the protocol. Class-wise IoU, Dice, confusion matrices and pixel accuracy are exported. Classes absent from both truth and prediction receive null scores instead of artificial perfect scores. The supplied implementation uses a fixed grouped split; automated cross-validation is not implemented.')
h('Binary development protocol')
p('The completed pilot is separate from this required four-class experiment. Its later runs use BCE plus Dice, AdamW, frozen batch-normalization statistics, paired rotations/flips and balanced replacement sampling. Native crack-centered training crops were added without changing the 20 validation crops. The final experiment added 240 eligible Dais patches and trained three further epochs; 198 distinct public patches were actually sampled. The crack threshold remained 0.5.')
page();h('Measured development results')
t=d.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Experiment','Crack Dice','Crack IoU']):c.text=v
for row in [('Pretrained balanced binary pilot','52.82%','35.89%'),('Additional augmented optimization','62.17%','45.11%'),('Additional UK training crops','68.45%','52.04%'),('Public masonry mixed training','73.25%','57.79%')]:
    for c,v in zip(t.add_row().cells,row):c.text=v
p('The final saved public-data checkpoint yields confusion counts TN 1269470, FP 4722, FN 12691 and TP 23837. An independent inference pass reproduced these counts. Pixel accuracy is 98.67%, but foreground Dice is the relevant overlap measure. The requested 80% Dice target has not been reached.')
p('These sequential experiments reuse the same development images and resume previous weights. The comparison is not an independent test, a controlled data-only ablation or a fair architecture ranking. Additional training and new data changed together. Building-level identity is unavailable for the UK sources, so scene-level leakage cannot be ruled out. Exact RGB and perceptual-hash screening found no specified matches between the added public subset and the UK sources, but cannot exclude every transformed crop or related scene.')
if (root/'public-validation-comparison.jpg').exists():
    page();h('Validation examples')
    d.add_picture(str(root/'public-validation-comparison.jpg'),width=Inches(4.4))
    p('Figure 1. Saved binary validation comparison. UK source imagery adapted from Storey and Snowball [1], CC BY 4.0. Development examples are not a held-out target-domain evaluation.')
page();h('Error analysis and target transfer')
p('A transfer illustration on Kurdish candidate 0650 shows remaining proposals along open joints and foreground details, while some crack-like plaster lines are missed. Fewer predicted regions after training do not prove better accuracy. No approved target mask exists for this example, so no numerical Kurdish accuracy is reported. Unmarked pixels remain unknown rather than certified healthy.')
p('Earlier automatic labeling attempts were unsuitable as ground truth. A CLIPSeg pseudo-label experiment produced very poor external crack overlap and visually misleading erosion predictions. The project preserves those negative findings to prevent their accidental reuse as approved labels. The binary model cannot generate validated erosion or vegetation classes.')
h('Software verification')
p('Unit checks cover confusion-matrix metrics, manifest validation and group leakage, ignored-pixel gradients, all-unknown loss handling, small-image padding and full-source pixel counting. A synthetic end-to-end check exercises all three architectures, training, checkpoint persistence and evaluation. Synthetic fixtures establish software behavior only and provide no scientific performance result.')
h('Deployment proposal')
p('A future HCECR inspection workflow would accept a site photograph and metadata, return class overlays with descriptive area fractions, and require an expert to approve or correct observations. Records should preserve original images, source/consent information, model version, predictions and review history. Low-quality or out-of-domain inputs should be flagged for manual inspection. Severity must use an agreed expert rubric and validated labels; the system must not issue automatic safety verdicts.')
p('The current deliverable is a local research workflow with saved visualizations. A smartphone service, organizational integration, professional validation, access controls and production monitoring have not been deployed. Before deployment, evaluate image privacy and access permissions, the consequences of false negatives, and performance on unseen buildings and camera conditions.')
page();h('Work required for final completion')
p('First, review candidate relevance and obtain at least 800 usable pixel-annotated images, replacing unsuitable candidates. Record annotator/reviewer identities and resolve ambiguous classes with the supervisor. Second, define structure groups and freeze the train/validation/test protocol. Third, produce target pixel class distributions, per-site statistics, color and edge analyses from the approved data.')
p('Fourth, train and compare all three architectures under the same protocol, including controlled preprocessing/loss comparisons where capacity allows. Report per-class scores, per-site errors and uncertainty across repeated seeds or grouped folds. Fifth, agree on severity labels and evaluate severity separately. Finally, revise the academic report with actual target results and prepare the final code/report submission. The deadline and severity rubric still require confirmation; no work has been submitted to the teacher.')
h('Reproducibility')
p('The repository contains source manifests, dataset hashes, experiment histories and environment locks. PUBLIC_MASONRY_RESULTS.md documents the completed binary run; FOUR_CLASS_RUN_GUIDE.md documents target training. At this historical snapshot, the public-data binary checkpoint was local at local-results/public-masonry/unet-diagnostic.pt and is not included in the repository. A fresh clone must reproduce the sequential training or obtain the checkpoint through an authorized distribution. Large raw data and model weights should not be assumed present.')
h('References')
p('[1] K. A. Storey and A. L. Snowball. Masonry crack images and masks. Zenodo record 18458458. https://zenodo.org/records/18458458. CC BY 4.0.')
p('[2] D. Dais, I. E. Bal, E. Smyrou and V. Sarhosis. Automatic crack classification and segmentation on masonry surfaces using convolutional neural networks and transfer learning. Automation in Construction 125, 103606, 2021. https://doi.org/10.1016/j.autcon.2021.103606. Code/data subset: https://github.com/dimitrisdais/crack_detection_CNN_masonry.')
p('[3] Unpublished project assignment brief, reviewed 18 October 2026. Institution, course identifiers and private source location are withheld in this public edition. The brief defines requirements; it is not evidence of completed experiments.')
for element in [d.styles.element,d.element]:
    for border in list(element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
# Public edition contains no author/recipient identity or machine metadata.
for table in d.tables:
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        line=OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:line.set(qn('w:'+key),value)
        borders.append(line)
    props.append(borders)
footer=d.sections[0].footer.paragraphs[0]
footer.text='Historical public-data research snapshot'
footer.style=d.styles['Normal']
for run in footer.runs:run.font.size=Pt(8)
d.core_properties.title='Historical citadel damage segmentation'
d.core_properties.subject='Historical public research progress report'
d.core_properties.author=''
d.core_properties.last_modified_by=''
d.core_properties.comments=''
d.core_properties.keywords=''
final=root/'ACADEMIC_REPORT_HISTORICAL_PUBLIC.docx'
if final.exists():raise FileExistsError('Use a fresh staging directory to preserve prior artifacts')
temporary=root/'report-before-metadata-scrub.docx'
d.save(temporary)
with zipfile.ZipFile(temporary) as original,zipfile.ZipFile(final,'w',zipfile.ZIP_DEFLATED) as cleaned:
    for item in original.infolist():
        data=original.read(item.filename)
        if item.filename=='docProps/core.xml':
            node=etree.fromstring(data)
            for child in list(node):
                if etree.QName(child).localname not in ('title','subject'):node.remove(child)
            data=etree.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
        elif item.filename=='docProps/app.xml':
            node=etree.fromstring(data)
            for child in list(node):
                if etree.QName(child).localname in ('Manager','Company','Template','Application','AppVersion','HyperlinkBase'):node.remove(child)
            data=etree.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
        elif item.filename.endswith('.xml'):
            node=etree.fromstring(data)
            for element in node.iter():
                for key in list(element.attrib):
                    if etree.QName(key).localname.startswith('rsid'):del element.attrib[key]
            data=etree.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
        cleaned.writestr(item,data)
temporary.unlink()
print(final.name)

